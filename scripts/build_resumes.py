"""Build DOCX + PDF resumes from the project's markdown resume files.

Usage:  python scripts/build_resumes.py
Outputs <name>.docx and <name>.pdf next to each source .md.

Sources:
  resume/master-resume.md
  applications/queue/*/resume.md
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent.parent
WIN_FONTS = Path("C:/Windows/Fonts")

SImap = {
    "\u2013": "-", "\u2014": "-", "\u2192": "->", "\u2022": "-",
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2713": "-", "\u26a0\ufe0f": "", "\u26a0": "", "\u2026": "...",
    "\u00b7": "|", "\u20b9": "INR ", "\u2011": "-", "\u00a0": " ",
}

def ascii_pdf(s: str) -> str:
    for k, v in SImap.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "ignore").decode("latin-1")


# ---------- markdown parse ----------
def parse(md: str):
    lines = md.splitlines()
    doc = {"name": "", "title": "", "contact": [], "sections": []}
    section = None
    seen_title = False
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("> "):            # editorial note -> skip
            continue
        if line.startswith("# "):
            doc["name"] = line[2:].strip()
            continue
        if line.startswith("## "):
            section = {"heading": line[3:].strip(), "items": []}
            doc["sections"].append(section)
            continue
        if re.match(r"#{3,6}\s", line) and section is not None:
            section["items"].append(("jobhead", line.lstrip("#").strip()))
            continue
        m_all_bold = re.fullmatch(r"\*\*(.+?)\*\*", line)
        if m_all_bold and section is None and not seen_title:
            doc["title"] = m_all_bold.group(1).strip()
            seen_title = True
            continue
        if section is None:                  # contact lines
            doc["contact"].append(line.strip())
            continue
        if line.startswith("- "):
            section["items"].append(("bullet", line[2:].strip()))
        elif line.startswith("**"):
            section["items"].append(("jobhead", line.strip()))
        else:
            section["items"].append(("para", line.strip()))
    return doc


def runs(text: str):
    """Split '**bold** normal' into [(txt, is_bold), ...]."""
    out = []
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            out.append((part, i % 2 == 1))
    return out


# ---------- DOCX ----------
def bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), "999999")
    bdr.append(b); pPr.append(bdr)


def build_docx(doc, out: Path):
    d = Document()
    st = d.styles["Normal"].font
    st.name = "Calibri"; st.size = Pt(10.5)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(0.5)
        s.left_margin = s.right_margin = Inches(0.7)

    h = d.add_paragraph()
    r = h.add_run(doc["name"]); r.bold = True; r.font.size = Pt(19)
    if doc["title"]:
        p = d.add_paragraph()
        rr = p.add_run(doc["title"]); rr.bold = True; rr.font.size = Pt(11.5)
        rr.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    for c in doc["contact"]:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(c).font.size = Pt(9.5)

    for sec in doc["sections"]:
        hp = d.add_paragraph()
        hp.paragraph_format.space_before = Pt(10); hp.paragraph_format.space_after = Pt(4)
        hr = hp.add_run(sec["heading"].upper()); hr.bold = True; hr.font.size = Pt(11)
        bottom_border(hp)
        for kind, text in sec["items"]:
            if kind == "bullet":
                p = d.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
            else:
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                if kind == "jobhead":
                    p.paragraph_format.space_before = Pt(6)
            for txt, b in runs(text):
                rn = p.add_run(txt); rn.bold = b
                if kind == "bullet":
                    rn.font.size = Pt(10)
    d.save(str(out))


# ---------- PDF ----------
class PDF(FPDF):
    def header(self):  # noqa
        pass


def build_pdf(doc, out: Path):
    pdf = PDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(18, 14, 18)
    pdf.add_page()
    try:
        pdf.add_font("Arial", "", str(WIN_FONTS / "arial.ttf"))
        pdf.add_font("Arial", "B", str(WIN_FONTS / "arialbd.ttf"))
        fam = "Arial"
    except Exception:
        fam = "Helvetica"
    conv = ascii_pdf  # keep PDF text in clean latin-1 for ATS parsers

    W = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.set_font(fam, "B", 18)
    pdf.multi_cell(W, 8, conv(doc["name"]))
    if doc["title"]:
        pdf.set_font(fam, "B", 11)
        pdf.set_text_color(60, 60, 60)
        pdf.multi_cell(W, 5.5, conv(doc["title"]))
        pdf.set_text_color(0, 0, 0)
    pdf.set_font(fam, "", 9)
    for c in doc["contact"]:
        pdf.multi_cell(W, 4.6, conv(c))
    pdf.ln(1.5)

    def write_rich(text, size, bullet=False):
        pdf.set_font(fam, "", size)
        x0 = pdf.l_margin + (5 if bullet else 0)
        if bullet:
            pdf.set_x(pdf.l_margin)
            pdf.cell(5, 4.8, conv("-"))
        pdf.set_x(x0)
        for txt, b in runs(text):
            pdf.set_font(fam, "B" if b else "", size)
            pdf.write(4.8, conv(txt))
        pdf.ln(5.2 if not bullet else 5.0)

    for sec in doc["sections"]:
        pdf.ln(2.5)
        pdf.set_font(fam, "B", 11)
        pdf.cell(0, 6, conv(sec["heading"].upper()), new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.set_draw_color(150, 150, 150)
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(1.5)
        for kind, text in sec["items"]:
            if kind == "jobhead":
                pdf.ln(1.5)
                write_rich(text, 10)
            elif kind == "bullet":
                write_rich(text, 9.5, bullet=True)
            else:
                write_rich(text, 9.5)
    pdf.output(str(out))


def main():
    targets = [ROOT / "resume" / "master-resume.md"]
    targets += sorted((ROOT / "applications" / "queue").glob("*/resume.md"))
    for md_path in targets:
        if not md_path.exists():
            continue
        doc = parse(md_path.read_text(encoding="utf-8"))
        stem = md_path.with_suffix("")
        build_docx(doc, stem.with_suffix(".docx"))
        build_pdf(doc, stem.with_suffix(".pdf"))
        print("built:", stem.with_suffix(".docx").relative_to(ROOT), "+ .pdf")


if __name__ == "__main__":
    sys.exit(main())
